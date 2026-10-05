import hashlib
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from feedback_pipeline.cli import run
from feedback_pipeline.config import resolve_period
from feedback_pipeline.report import build_report_plan, validate_analysis
from feedback_pipeline.sources import load_sources

ROOT = Path(__file__).resolve().parents[1]


class LocalPipelineTests(unittest.TestCase):
    def test_complete_demo_counts_provenance_and_input_immutability(self):
        import yaml
        cfg = yaml.safe_load((ROOT / "config.example.yaml").read_text())
        before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT / "examples").iterdir()}
        with tempfile.TemporaryDirectory() as temp:
            cfg["sources"]["files"] = [{**s, "path": str(ROOT / s["path"])} for s in cfg["sources"]["files"]]
            cfg["analysis"]["rules_file"] = str(ROOT / cfg["analysis"]["rules_file"])
            cfg["output"]["directory"] = str(Path(temp) / "outputs")
            config = Path(temp) / "demo.yaml"
            config.write_text(yaml.safe_dump(cfg, allow_unicode=True))
            result = run(config)
            self.assertEqual((result["records"], result["verified_period_records"]), (25, 24))
            out = Path(result["output_directory"])
            quality = json.loads((out / "quality.json").read_text())
            self.assertEqual((quality["duplicates"], quality["out_of_period"]), (2, 4))
            analysis = json.loads((out / "analysis.json").read_text())
            validate_analysis(analysis)
            self.assertEqual(analysis["totals"]["rated_store_records"], 12)
            self.assertEqual(analysis["totals"]["negative"], 10)
            self.assertEqual(analysis["totals"]["negative_rate"], 83.3)
            self.assertTrue(any(len(r["locations"]) == 2 for r in analysis["records"].values()))
            self.assertTrue(any(l["row_end"] > l["row"] for r in analysis["records"].values() for l in r["locations"]))
            workbook = load_workbook(out / "details.xlsx", read_only=True)
            self.assertEqual(workbook["Records"].max_row, 26)
            self.assertEqual(workbook["Summary"]["B6"].value, 0.833)
            workbook.close()
            state = json.loads((out / "run.json").read_text())
            self.assertTrue(all(hashlib.sha256((out / n).read_bytes()).hexdigest() == h for n, h in state["files"].items()))
        self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in before})

    def test_excel_rejects_oversized_text_without_silent_truncation(self):
        from feedback_pipeline.analysis import analyze_records
        from feedback_pipeline.config import load_yaml
        from feedback_pipeline.exporters import write_excel
        r = {"id": "large", "source": "external", "platform": "forum", "date": "2025-02-02", "period_start": "2025-02-02", "period_end": "2025-02-02", "rating": None,
             "title": "", "text": "长" * 32768, "file": "input.csv", "sheet": "", "row": 2}
        payload = {"records": [r], "period": {"start_date": "2025-02-01", "end_date": "2025-02-07"}, "files": [], "matched_sheets": [], "warnings": []}
        plan = build_report_plan(analyze_records(payload, load_yaml(ROOT / "config/rules.generic.yaml")), {})
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / "details.xlsx"
            with self.assertRaises(ValueError):
                write_excel(plan, {"files": [], "excluded": []}, out)
            self.assertFalse(out.exists())

    def test_xlsx_header_detection_and_mapping(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            wb = Workbook()
            sheet = wb.active
            sheet.append(["Export note"])
            sheet.append([None, "日期", "内容", "评分"])
            sheet.append([None, datetime(2025, 2, 2), "登录失败，验证码收不到", 2])
            source = root / "input.xlsx"
            wb.save(source)
            payload = load_sources([{"id": "workbook", "path": "input.xlsx", "kind": "ios", "header_row": "auto"}], "2025-02-01", "2025-02-07", root)
            self.assertEqual(len(payload["records"]), 1)
            self.assertEqual(payload["records"][0]["row"], 3)
            self.assertEqual(payload["records"][0]["rating"], 2)

    def test_invalid_date_and_truncated_csv_fail(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "bad.csv"
            entry = [{"id": "bad", "path": "bad.csv", "kind": "ios"}]
            source.write_text("date,text,rating\n2025-02-31,登录失败,2\n")
            with self.assertRaises(ValueError):
                load_sources(entry, "2025-02-01", "2025-02-07", root)
            source.write_text("date,text,rating\n2025-02-02,登录失败\n")
            with self.assertRaises(ValueError):
                load_sources(entry, "2025-02-01", "2025-02-07", root)

    def test_customer_role_extraction_keeps_full_dialogue_dedup_identity(self):
        import csv
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            path = root / "support.csv"
            with path.open("w", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["date", "text"])
                writer.writerow(["2025-02-02", "用户：登录失败，验证码收不到。\n客服：请重新登录。"])
                writer.writerow(["2025-02-02", "用户：登录失败，验证码收不到。\n客服：请检查网络。"])
                writer.writerow(["2025-02-02", "客服：正在处理。系统：已结束。"])
            payload = load_sources([{"id": "support", "path": "support.csv", "kind": "customer_service"}], "2025-02-01", "2025-02-07", root)
            self.assertEqual(len(payload["records"]), 2)
            self.assertTrue(all(r["text"] == "登录失败，验证码收不到。" for r in payload["records"]))
            self.assertTrue(all(r["rating"] is None for r in payload["records"]))
            self.assertEqual(payload["quality"]["files"][0]["duplicates"], 0)

    def test_redaction_never_changes_ids_and_excel_formula_text(self):
        from feedback_pipeline.analysis import analyze_records
        from feedback_pipeline.config import load_yaml
        from feedback_pipeline.exporters import write_excel
        contact = "1" + "38" + "0" * 8
        rid = "record:" + contact + "abc"
        record = {"id": rid, "source": "ios", "platform": "iOS", "date": "2025-02-02", "period_start": "2025-02-02", "period_end": "2025-02-02", "rating": 1,
                  "title": "", "text": "=登录失败，联系方式 " + contact, "file": "input.csv", "sheet": "", "row": 2, "raw": {}}
        payload = {"records": [record], "period": {"start_date": "2025-02-01", "end_date": "2025-02-07"}, "files": [], "matched_sheets": [], "warnings": []}
        analysis = analyze_records(payload, load_yaml(ROOT / "config/rules.generic.yaml"))
        plan = build_report_plan(analysis, {"redact_contacts": True})
        self.assertEqual(set(plan["records"]), {rid})
        self.assertNotIn(contact, plan["records"][rid]["text"])
        validate_analysis(plan)
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp) / "details.xlsx"
            write_excel(plan, {"files": [], "excluded": []}, out)
            workbook = load_workbook(out)
            self.assertEqual(workbook["Records"]["I2"].data_type, "s")
            workbook.close()

    def test_dry_run_writes_nothing_and_default_period_crosses_year(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = ROOT / "config.example.yaml"
            config = root / "config.yaml"
            config.write_text(source.read_text().replace("config/rules.generic.yaml", str(ROOT / "config/rules.generic.yaml")))
            before = list(root.iterdir())
            self.assertTrue(run(config, dry_run=True)["dry_run"])
            self.assertEqual(before, list(root.iterdir()))
        self.assertEqual(resolve_period({}, now=datetime(2025, 1, 3, tzinfo=ZoneInfo("Asia/Shanghai"))), ("2024-12-27", "2025-01-02"))
        with self.assertRaises(ValueError):
            resolve_period({}, "2025-02-01", None)

    def test_contact_redaction_covers_generic_and_media_export_evidence(self):
        import csv
        import yaml
        from feedback_pipeline.analysis import analyze_dynamic_records, analyze_records
        from feedback_pipeline.config import load_yaml

        phone = "1" + "38" + "0" * 8
        email = "fixture" + "@" + "example.invalid"
        contact = f"邮箱 {email}，电话 {phone}"
        for mode, analyzer in [("generic", analyze_records), ("media", analyze_dynamic_records)]:
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                source = root / "input.csv"
                with source.open("w", newline="", encoding="utf-8") as stream:
                    writer = csv.writer(stream)
                    writer.writerow(["source", "platform", "date", "rating", "title", "text"])
                    writer.writerow(["ios", "demo-store", "2025-02-02", 1, "登录失败",
                                     "登录失败，视频无法播放，" + contact])
                    writer.writerow(["customer_service", "support", "2025-02-02", "", "",
                                     "用户：视频无法播放，" + contact + "。客服：请您提供截图。"])
                    writer.writerow(["external", "community", "2025-02-02", "", "", "仅供联系，" + contact])
                source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
                entries = [{"id": "synthetic", "path": source.name}]
                payload = load_sources(entries, "2025-02-01", "2025-02-07", root)
                rules_path = ROOT / f"config/rules.{mode}.yaml"
                original = analyzer(payload, load_yaml(rules_path))
                # These real pipeline shapes used to bypass report redaction.
                if mode == "generic":
                    self.assertTrue(any(phone in q["full_text"] for i in original["issues"] for q in i["quotes"]))
                else:
                    self.assertTrue(any(phone in text for i in original["issues"] for text in i["evidence_samples"]))
                    self.assertTrue(any(phone in text for i in original["issues"] for text in i["audit_texts"].values()))
                cfg = {"app": {"name": "Synthetic demo"},
                       "period": {"start": "2025-02-01", "end": "2025-02-07"},
                       "sources": {"files": entries},
                       "analysis": {"mode": mode, "rules_file": str(rules_path)},
                       "output": {"directory": "outputs", "formats": ["html", "markdown", "xlsx"]}}
                config = root / "config.yaml"
                config.write_text(yaml.safe_dump(cfg), encoding="utf-8")
                result = run(config)  # Default privacy must protect all generated formats.
                output = Path(result["output_directory"])
                for name in ("analysis.json", "report.html", "report.md"):
                    text = (output / name).read_text(encoding="utf-8")
                    self.assertNotIn(phone, text)
                    self.assertNotIn(email, text)
                exported = json.loads((output / "analysis.json").read_text(encoding="utf-8"))
                validate_analysis(exported)
                self.assertEqual(set(exported["records"]), set(original["records"]))
                self.assertEqual([i["issue_key"] for i in exported["issues"]], [i["issue_key"] for i in original["issues"]])
                self.assertEqual([i["record_ids"] for i in exported["issues"]], [i["record_ids"] for i in original["issues"]])
                self.assertEqual(exported["sources"]["files"], original["sources"]["files"])
                workbook = load_workbook(output / "details.xlsx", read_only=True)
                try:
                    for sheet in workbook:
                        for row in sheet.iter_rows(values_only=True):
                            for value in row:
                                if isinstance(value, str):
                                    self.assertNotIn(phone, value)
                                    self.assertNotIn(email, value)
                finally:
                    workbook.close()
                self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_hash)
                unredacted = build_report_plan(original, {"redact_contacts": False})
                self.assertTrue(any(phone in r["text"] for r in unredacted["records"].values()))

    def test_redacted_truncated_quote_remains_a_source_substring(self):
        from feedback_pipeline.analysis import analyze_records
        from feedback_pipeline.config import load_yaml

        phone = "1" + "38" + "0" * 8
        email = "fixture" + "@" + "example.invalid"
        for contact in (phone, email):
            text = "登录失败，联系方式 " + contact
            rules = load_yaml(ROOT / "config/rules.generic.yaml")
            rules["rules"]["max_quote_chars"] = len(text) - 3
            record = {"id": "stable-id", "source": "ios", "platform": "demo", "date": "2025-02-02",
                      "rating": 1, "title": "", "text": text, "file": "input.csv", "sheet": "", "row": 2}
            payload = {"records": [record], "period": {"start_date": "2025-02-01", "end_date": "2025-02-07"},
                       "files": [], "matched_sheets": [], "warnings": []}
            analysis = analyze_records(payload, rules)
            plan = build_report_plan(analysis, {})
            validate_analysis(plan)
            for issue in plan["issues"]:
                for quote in issue["quotes"]:
                    self.assertNotIn(contact[:len(contact) - 3], quote["quote"])
                    self.assertIn(quote["quote"], plan["records"]["stable-id"]["text"])

    def test_qimai_runtime_configs_are_unique_and_cleaned_on_success_and_failure(self):
        import subprocess
        from unittest.mock import patch
        from feedback_pipeline.cli import fetch_qimai

        paths = []
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = {"app": {"name": "Synthetic demo", "ios_app_id": "1"},
                      "sources": {"qimai": {"platforms": ["ios"]}}}
            for returncode in (0, 1):
                def execute(command, **kwargs):
                    path = Path(command[command.index("--config") + 1])
                    self.assertTrue(path.is_file())
                    self.assertEqual(json.loads(path.read_text())["app"]["ios"]["appid"], "1")
                    paths.append(path)
                    if returncode == 0:
                        state = {"ios": {"download": {"path": str(root / "synthetic.csv")}}, "android": []}
                        (root / "outputs/qimai/qimai_fetch_2025-02-01_2025-02-07.json").write_text(json.dumps(state))
                    return subprocess.CompletedProcess(command, returncode)
                with patch("feedback_pipeline.cli.ROOT", root), patch("feedback_pipeline.cli.shutil.which", return_value="node"), \
                     patch("feedback_pipeline.cli.subprocess.run", side_effect=execute):
                    if returncode:
                        with self.assertRaises(ValueError):
                            fetch_qimai(config, "2025-02-01", "2025-02-07")
                    else:
                        entries = fetch_qimai(config, "2025-02-01", "2025-02-07")
                        self.assertEqual(entries[0]["kind"], "ios")
                self.assertFalse(paths[-1].exists())
                self.assertFalse(paths[-1].parent.exists())
        self.assertNotEqual(paths[0], paths[1])


if __name__ == "__main__":
    unittest.main()
