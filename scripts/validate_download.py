"""Validate that a browser download is complete and readable."""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl


def check_dates(rows, start, end):
    if not start:
        return {}
    date_index = None
    count = 0
    for row in rows:
        if date_index is None:
            date_index = next((i for i, cell in enumerate(row) if str(cell).strip() in ("发表时间", "评论时间", "日期")), None)
            continue
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        value = str(row[date_index]) if date_index < len(row) else ""
        match = re.search(r"(20\d{2})[-/](\d{1,2})[-/](\d{1,2})", value)
        if not match:
            raise ValueError("unrecognized comment date")
        date = "{}-{:02d}-{:02d}".format(match[1], int(match[2]), int(match[3]))
        if not start <= date <= end:
            raise ValueError("comment outside requested period: " + date)
        count += 1
    if date_index is None:
        raise ValueError("comment date header missing")
    return {"period_checked": True, "comment_rows": count}


def validate_xlsx(path: Path, start=None, end=None) -> dict:
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"expected .xlsx, got {path.suffix}")
    if not zipfile.is_zipfile(path):
        raise ValueError("not a readable ZIP package")
    with zipfile.ZipFile(path) as package:
        broken = package.testzip()
        if broken:
            raise ValueError(f"corrupt ZIP member: {broken}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = workbook.sheetnames
        sheet = workbook[sheets[0]]
        sheet.reset_dimensions()
        period = check_dates(sheet.iter_rows(values_only=True), start, end)
    finally:
        workbook.close()
    return {"kind": "xlsx", "sheets": sheets, **period}


def validate_csv(path: Path, start=None, end=None) -> dict:
    raw = path.read_bytes()
    decoded = None
    encoding = None
    for candidate in ("utf-8-sig", "gb18030"):
        try:
            decoded = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValueError("CSV encoding is not readable")
    rows = list(csv.reader(io.StringIO(decoded, newline=""), strict=True))
    if not rows or not any(cell.strip() for cell in rows[0]):
        raise ValueError("CSV header is missing")
    if any(len(row) != len(rows[0]) for row in rows if row):
        raise ValueError("CSV column count mismatch")
    return {"kind": "csv", "encoding": encoding, "rows": len(rows), **check_dates(rows, start, end)}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("--start-date")
    parser.add_argument("--end-date")
    args = parser.parse_args()
    path = args.path.resolve()
    try:
        if not path.is_file():
            raise ValueError("file does not exist")
        if path.name.endswith((".crdownload", ".part", ".tmp")):
            raise ValueError("temporary download suffix remains")
        size = path.stat().st_size
        if size <= 0:
            raise ValueError("file is empty")
        if path.suffix.lower() == ".xlsx":
            details = validate_xlsx(path, args.start_date, args.end_date)
        elif path.suffix.lower() == ".csv":
            details = validate_csv(path, args.start_date, args.end_date)
        else:
            details = {"kind": path.suffix.lower().lstrip(".")}
        print(json.dumps({"ok": True, "path": str(path), "size": size, **details}, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "path": str(path), "error": str(exc)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
