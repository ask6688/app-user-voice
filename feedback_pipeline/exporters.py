"""Render the shared report plan locally, without any account or publishing API."""
from __future__ import annotations

import json
import re
from datetime import date
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .report import full_text, html_text, provenance

# The original local report's readable table layout is retained and adapted.
STYLE = """<style>
body{font-family:Arial,'PingFang SC',sans-serif;color:#222;line-height:1.6;max-width:1100px;margin:36px auto;padding:0 20px}
h1{font-size:25px}h2{font-size:19px;margin-top:32px}h3{font-size:16px}
blockquote{border-left:3px solid #b5c4d7;padding-left:14px;color:#536173;margin-left:0}
table{border-collapse:collapse;width:100%;table-layout:fixed;font-size:13px}
td,th{border:1px solid #e2e4e8;padding:8px;vertical-align:top;text-align:left;overflow-wrap:anywhere}
th{background:#f5f6f8}code{font-size:12px;background:#f2f2f2}a{color:#315d93}
.warning{background:#fff7e6;padding:12px;border-left:3px solid #b17a24}.meta{color:#657080}
@media print{body{margin:0;max-width:none}h2,h3{break-after:avoid}tr{break-inside:avoid}}
</style>"""


def rate_text(totals: dict) -> str:
    denominator = totals.get("rated_store_records", totals["store_records"])
    value = totals.get("negative_rate")
    return "不适用（没有有效商店评分）" if value is None else f"{value}%（{totals['negative']}/{denominator}）"


def _warnings(plan: dict) -> list[str]:
    return plan.get("metadata", {}).get("warnings", plan.get("warnings", []))


def _title(plan: dict, app: dict) -> str:
    return f"{app['name']} 用户声音周报｜{plan['period']['start_date']}—{plan['period']['end_date']}"


def write_html(plan: dict, app: dict, path: Path) -> None:
    totals = plan["totals"]
    parts = [f"<h1>{escape(_title(plan, app))}</h1>",
             f"<p class='meta'>{escape(app.get('company', '')) + ' · ' if app.get('company') else ''}基于本次输入文件的规则分析</p>",
             "<h2>本期概览</h2><table><tr><th>反馈总量</th><th>有效商店评分</th><th>商店差评率</th><th>问题数</th><th>未分类</th></tr>",
             f"<tr><td>{totals['records']}</td><td>{totals.get('rated_store_records', totals['store_records'])}</td><td>{escape(rate_text(totals))}</td><td>{len(plan['issues'])}</td><td>{totals['unclassified']}</td></tr></table>",
             "<p>差评仅指配置阈值内的商店评分，外部舆情和客服反馈不参与分母。来源缺日期的记录保留，但无法证明属于本周期。</p>"]
    if _warnings(plan):
        parts.append("<div class='warning'><strong>数据检查</strong><ul>" + "".join(f"<li>{html_text(w)}</li>" for w in _warnings(plan)) + "</ul></div>")
    parts.append("<h2>来源与评分分布</h2><table><tr><th>来源文件</th><th>读取行数</th><th>保留记录</th><th>周期外</th><th>重复</th><th>缺日期</th></tr>")
    for f in plan["sources"]["files"]:
        parts.append("<tr>" + "".join(f"<td>{html_text(f.get(k, ''))}</td>" for k in ["file", "rows_read", "records_kept", "out_of_period", "duplicates", "missing_date"]) + "</tr>")
    parts.append("</table><p>商店评分：" + "；".join(f"{star} 星 {n} 条" for star, n in totals["rating_distribution"].items()) + "</p><h2>问题与代表原声</h2>")
    for category in plan["categories"]:
        parts.append(f"<h3>{html_text(category['module_title'])}（{category['record_count']} 条反馈）</h3>")
        for issue in [i for i in plan["issues"] if i["category_key"] == category["key"]]:
            parts.append(f"<p><strong>{html_text(issue['code'])} · {html_text(issue['name'])}</strong>：{issue['record_count']} 条，商店差评 {issue['negative_count']} 条</p>")
            for q in issue.get("quotes", []):
                r = plan["records"][q["record_id"]]
                parts.append(f"<blockquote>{html_text(q['quote'])}<br><small>{html_text(r['platform'])}｜{html_text(provenance(r))}</small></blockquote>")
    if not plan["issues"]:
        parts.append("<p>本期没有匹配到已配置的问题主题。请查看未分类反馈或调整规则。</p>")
    parts.append("<h2>原文核对表</h2><p>每个问题展示有限原文；问题计数使用完整记录集合，全部反馈保存在明细文件中。</p>")
    for issue in plan["issues"]:
        parts.append(f"<h3>{html_text(issue['code'])} · {html_text(issue['name'])}</h3><table><tr><th>原文</th><th>出处</th></tr>")
        for rid in issue["audit_record_ids"]:
            r = plan["records"][rid]
            link = f"<br><a href='{escape(r['url'], quote=True)}' rel='noreferrer'>原帖</a>" if r.get("url") else ""
            parts.append(f"<tr><td>{html_text(full_text(r))}</td><td>{html_text(r['platform'])}｜{html_text(r.get('date') or '日期缺失')}｜{html_text(provenance(r))}{link}<br><code>{html_text(rid)}</code></td></tr>")
        parts.append("</table>")
    parts.append("<h2>未分类反馈</h2><table><tr><th>内容</th><th>出处</th></tr>")
    for u in plan["unclassified"]:
        r = plan["records"][u["id"]]
        parts.append(f"<tr><td>{html_text(r['text'])}</td><td>{html_text(provenance(r))}</td></tr>")
    parts.append("</table><p class='meta'>自动报告供人工复核。多标签反馈可进入多个问题，问题数量之和不等于独立反馈总量。联系方式按配置遮盖，原始输入文件未被修改。</p>")
    path.write_text("<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>" + escape(_title(plan, app)) + "</title>" + STYLE + "</head><body>" + "".join(parts) + "</body></html>", encoding="utf-8")


def write_markdown(plan: dict, app: dict, path: Path) -> None:
    def md(text):
        return str(text).replace("\\", "\\\\").replace("[", "\\[").replace("]", "\\]").replace("<", "&lt;").replace("|", "\\|")
    lines = [f"# {md(_title(plan, app))}", "", f"反馈总量：{plan['totals']['records']}；商店差评率：{rate_text(plan['totals'])}；未分类：{plan['totals']['unclassified']}。", "",
             "差评率仅使用有有效评分的商店反馈，外部舆情和客服不参与分母。", ""]
    if _warnings(plan):
        lines += ["## 数据检查", ""] + [f"- {md(w)}" for w in _warnings(plan)] + [""]
    for category in plan["categories"]:
        lines += [f"## {md(category['module_title'])}", ""]
        for issue in [i for i in plan["issues"] if i["category_key"] == category["key"]]:
            lines += [f"### {md(issue['code'])} · {md(issue['name'])}", "", f"反馈 {issue['record_count']} 条；商店差评 {issue['negative_count']} 条。", ""]
            for q in issue.get("quotes", []):
                r = plan["records"][q["record_id"]]
                lines += ["> " + md(q["quote"]).replace("\n", "\n> "), "", f"出处：{md(provenance(r))}；记录 `{r['id']}`。", ""]
            lines += ["原文核对：", ""]
            for rid in issue["audit_record_ids"]:
                r = plan["records"][rid]
                lines += [f"- `{rid}` · {md(provenance(r))}", "", "  > " + md(full_text(r)).replace("\n", "\n  > "), ""]
    lines += ["## 未分类反馈", ""]
    for u in plan["unclassified"]:
        r = plan["records"][u["id"]]
        lines += [f"- {md(r['text'])}（{md(provenance(r))}）", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_excel(plan: dict, quality: dict, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    def table(name, headers, rows, widths):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            for value in row:
                if isinstance(value, str) and (len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)):
                    raise ValueError("Excel 文本超出支持范围，请选择 HTML/Markdown 输出")
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315D73")
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    if len(cell.value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", cell.value):
                        raise ValueError("Excel 文本超出支持范围，请选择 HTML/Markdown 输出")
                    cell.data_type = "s"  # Keep user text beginning with '=' as text.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        return sheet
    t = plan["totals"]
    table("Summary", ["指标", "数值"], [["反馈总量", t["records"]], ["商店反馈", t["store_records"]],
          ["有效商店评分", t.get("rated_store_records", t["store_records"])], ["商店差评", t["negative"]],
          ["商店差评率", None if t["negative_rate"] is None else t["negative_rate"] / 100],
          ["问题数", len(plan["issues"])], ["未分类", t["unclassified"]]], [28, 24])
    workbook["Summary"]["B6"].number_format = "0.0%"
    table("Issues", ["编号", "主题", "问题", "反馈数", "商店差评", "完整记录ID", "代表原声"],
          [[i["code"], i["category_name"], i["name"], i["record_count"], i["negative_count"],
            "\n".join(i["record_ids"]), "\n".join(q["quote"] for q in i["quotes"])] for i in plan["issues"]],
          [12, 22, 40, 12, 12, 42, 70])
    issue_codes = {}
    for i in plan["issues"]:
        for rid in i["record_ids"]:
            issue_codes.setdefault(rid, []).append(i["code"])
    table("Records", ["记录ID", "来源", "渠道", "日期", "期间开始", "期间结束", "评分", "标题", "正文", "问题编号", "出处", "原帖"],
          [[r["id"], r["source"], r["platform"], date.fromisoformat(r["date"]) if r.get("date") else None,
            r.get("period_start"), r.get("period_end"), r["rating"], r["title"], r["text"],
            ",".join(issue_codes.get(r["id"], [])), provenance(r), r.get("url", "")] for r in plan["records"].values()],
          [36, 18, 18, 14, 14, 14, 9, 25, 70, 18, 55, 36])
    for cell in workbook["Records"]["D"][1:]:
        cell.number_format = "yyyy-mm-dd"
    table("Sources", ["来源ID", "文件", "SHA256", "读取行", "保留记录", "周期外", "重复", "缺日期", "缺评分"],
          [[f.get(k) for k in ["source_id", "file", "sha256", "rows_read", "records_kept", "out_of_period", "duplicates", "missing_date", "missing_rating"]] for f in quality["files"]],
          [22, 30, 68, 12, 12, 12, 12, 12, 12])
    table("Excluded", ["原因", "文件", "工作表", "起始行", "结束行", "重复记录ID"],
          [[e.get(k) for k in ["reason", "file", "sheet", "row", "row_end", "record_id"]] for e in quality["excluded"]],
          [24, 30, 22, 12, 12, 36])
    table("Warnings", ["数据告警"], [[w] for w in _warnings(plan)], [100])
    workbook.properties.creator = "App User Voice"
    workbook.properties.lastModifiedBy = "App User Voice"
    workbook.save(path)


def export_report(plan: dict, app: dict, quality: dict, directory: Path, formats: list[str]) -> list[str]:
    directory.mkdir(parents=True, exist_ok=True)
    outputs = []
    for fmt in formats:
        name = {"html": "report.html", "markdown": "report.md", "xlsx": "details.xlsx"}[fmt]
        target = directory / name
        if fmt == "html":
            write_html(plan, app, target)
        elif fmt == "markdown":
            write_markdown(plan, app, target)
        else:
            write_excel(plan, quality, target)
        outputs.append(name)
    for name, value in [("analysis.json", plan), ("quality.json", quality)]:
        (directory / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
        outputs.append(name)
    return outputs
