"""Import files without rewriting them; keep row provenance and quality counts."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .customer_service import ROLES, meaningful, plain_message, user_messages

KINDS = {"ios", "android", "external", "customer_service"}
ALIASES = {
    "source": ["source", "来源类型"],
    "platform": ["platform", "平台", "渠道"],
    "date": ["date", "发表时间", "评论时间", "反馈时间", "日期", "时间"],
    "rating": ["rating", "评分", "星级", "评级"],
    "title": ["title", "标题"],
    "text": ["text", "内容", "评论内容", "反馈内容", "用户内容", "反馈详情", "评价"],
    "url": ["url", "地址", "原文链接"],
    "review_id": ["review_id", "评论ID"],
    "period_start": ["period_start"], "period_end": ["period_end"],
}


def _string(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def parse_date(value) -> str | None:
    text = _string(value)
    if not text:
        return None
    match = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?(?:[ T].*)?", text)
    if not match:
        raise ValueError("日期无法识别，请使用包含年份的日期")
    return date(*(int(x) for x in match.groups())).isoformat()


def parse_rating(value) -> int | None:
    text = _string(value)
    if not text or text in {"-", "无", "N/A"}:
        return None
    if not re.fullmatch(r"[1-5](?:\.0)?(?:星|分)?", text):
        raise ValueError("评分必须为空或 1–5 的整数")
    return int(text[0])


def _tables(path: Path, entry: dict):
    if path.suffix.lower() in {".csv", ".tsv"}:
        raw = path.read_bytes()
        encoding = entry.get("encoding")
        if encoding:
            text = raw.decode(encoding)
        else:
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError(f"文件编码无法识别：{path.name}")
        delimiter = entry.get("delimiter", "\t" if path.suffix.lower() == ".tsv" else ",")
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        rows = []
        while True:
            first = reader.line_num + 1
            try:
                row = next(reader)
            except StopIteration:
                break
            rows.append((row, first, reader.line_num))
        yield "", rows, encoding
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            names = entry.get("sheets", workbook.sheetnames)
            if isinstance(names, str):
                names = [names]
            for name in names:
                if name not in workbook:
                    raise ValueError(f"找不到工作表：{path.name}/{name}")
                yield name, [(list(row), i, i) for i, row in
                             enumerate(workbook[name].iter_rows(values_only=True), 1)], "xlsx"
        finally:
            workbook.close()
    else:
        raise ValueError(f"暂不支持该文件格式：{path.suffix}")


def _mapped(headers: list[str], columns: dict) -> dict:
    mapping = {}
    for field, aliases in ALIASES.items():
        if field in columns:
            name = columns[field]
            if name not in headers:
                raise ValueError(f"配置字段不存在：{name}")
        else:
            name = next((key for key in aliases if key in headers), None)
        if name is not None:
            mapping[field] = headers.index(name)
    if "text" not in mapping:
        raise ValueError("找不到正文列，请在 columns.text 配置实际表头")
    return mapping


def load_sources(entries: list[dict], start: str, end: str, base_dir: Path) -> dict:
    records, by_identity, files, excluded, warnings = [], {}, [], [], []
    for entry in entries:
        if entry.get("enabled", True) is False:
            continue
        path = (base_dir / entry["path"]).resolve()
        if not path.is_file():
            raise ValueError(f"输入文件不存在：{entry['id']}")
        summary = {"source_id": entry["id"], "file": path.name,
                   "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                   "rows_read": 0, "blank_rows": 0, "out_of_period": 0,
                   "duplicates": 0, "missing_date": 0, "missing_rating": 0,
                   "records_kept": 0, "transformed_dialogues": 0, "sheets": []}
        for sheet, table, encoding in _tables(path, entry):
            summary["encoding"] = encoding
            summary["sheets"].append(sheet)
            header_row = entry.get("header_row", 1)
            if header_row == "auto":
                requested_text = entry.get("columns", {}).get("text")
                candidates = [requested_text] if requested_text else ALIASES["text"]
                offset = next((i for i, (cells, _, _) in enumerate(table[:30])
                               if any(_string(cell) in candidates for cell in cells)), -1)
            else:
                offset = int(header_row) - 1
            if offset < 0 or offset >= len(table):
                raise ValueError(f"表头行不存在：{path.name}")
            headers = [_string(cell) for cell in table[offset][0]]
            nonempty = [h for h in headers if h]
            if not nonempty or len(nonempty) != len(set(nonempty)):
                raise ValueError(f"表头为空或重名：{path.name}")
            mapping = _mapped(headers, entry.get("columns", {}))
            for values, first, last in table[offset + 1:]:
                summary["rows_read"] += 1
                location = {"source_id": entry["id"], "file": path.name,
                            "sheet": sheet, "row": first, "row_end": last}
                if not any(_string(cell) for cell in values):
                    summary["blank_rows"] += 1
                    continue
                if len(values) > len(headers) or (path.suffix.lower() != ".xlsx" and len(values) != len(headers)):
                    raise ValueError(f"列数不一致：{path.name} 第 {first} 行")
                values += [None] * (len(headers) - len(values))
                data = {key: values[index] for key, index in mapping.items()}
                kind = _string(data.get("source")) or entry.get("kind", "external")
                if kind not in KINDS:
                    raise ValueError(f"来源类型不支持：{path.name} 第 {first} 行")
                platform = _string(data.get("platform")) or entry.get("platform", kind)
                title, text = _string(data.get("title")), _string(data.get("text"))
                identity_text, transformation = text, "none"
                if kind == "customer_service":
                    plain = plain_message(text)
                    role = ROLES.match(plain)
                    if role:
                        segments = [plain_message(m) for m in user_messages(plain)]
                        text = "；".join(m for m in segments if meaningful(m))
                        transformation = "customer_user_segments"
                        summary["transformed_dialogues"] += 1
                        if not text:
                            summary["blank_rows"] += 1
                            excluded.append({**location, "reason": "no_meaningful_user_text"})
                            continue
                if not text and not title:
                    summary["blank_rows"] += 1
                    continue
                try:
                    point = parse_date(data.get("date"))
                    a = parse_date(data.get("period_start")) or point
                    b = parse_date(data.get("period_end")) or point
                    rating = parse_rating(data.get("rating")) if kind in {"ios", "android"} else None
                except ValueError as exc:
                    raise ValueError(f"{path.name} 第 {first} 行：{exc}") from exc
                if (a is None) != (b is None) or (a and b and a > b):
                    raise ValueError(f"时间范围无效：{path.name} 第 {first} 行")
                if a and b and not (start <= a <= b <= end):
                    summary["out_of_period"] += 1
                    excluded.append({**location, "reason": "out_of_period"})
                    continue
                if not a:
                    summary["missing_date"] += 1
                if rating is None and kind in {"ios", "android"}:
                    summary["missing_rating"] += 1
                review_id = _string(data.get("review_id"))
                identity = json.dumps([kind, platform, a, b, rating, title, identity_text, review_id], ensure_ascii=False)
                rid = "record:" + hashlib.sha256(identity.encode()).hexdigest()[:24]
                if identity in by_identity:
                    original = by_identity[identity]
                    original["locations"].append(location)
                    summary["duplicates"] += 1
                    excluded.append({**location, "reason": "exact_duplicate", "record_id": original["id"]})
                    continue
                record = {"id": rid, "source": kind, "platform": platform,
                          "date": point or a, "period_start": a, "period_end": b,
                          "rating": rating, "title": title, "text": text,
                          "url": _string(data.get("url")), "file": path.name,
                          "sheet": sheet, "row": first, "locations": [location],
                          "text_transformation": transformation,
                          "raw": {"标题": title, "内容": text}}
                records.append(record)
                by_identity[identity] = record
                summary["records_kept"] += 1
        if summary["missing_date"]:
            warnings.append(f"{entry['id']}：{summary['missing_date']} 条反馈缺少日期，无法确认所属周期，保留并标记")
        if summary["missing_rating"]:
            warnings.append(f"{entry['id']}：{summary['missing_rating']} 条商店反馈缺少评分，不进入差评率分母")
        files.append(summary)
    if not files:
        raise ValueError("没有启用的文件来源")
    return {"period": {"start_date": start, "end_date": end}, "records": records,
            "files": files, "matched_sheets": [{"sheet": sheet, "file": f["file"]}
                       for f in files for sheet in f["sheets"]], "warnings": warnings,
            "quality": {"files": files, "records": len(records), "excluded": excluded,
                        "out_of_period": sum(f["out_of_period"] for f in files),
                        "duplicates": sum(f["duplicates"] for f in files)}}
